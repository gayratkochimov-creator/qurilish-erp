"""Excel'dagi material ro'yxatini obyektning UMUMIY LIMITIGA QORALAMA sifatida kiritadi.

Kutilgan jadval (sarlavha qatori avtomatik topiladi, ustun nomlari bo'yicha):
  № | <bo'lim/guruh> | наименование | ед/из | число | цена | сумма нал | сумма переч | премечание
- 2-ustun (guruh nomi) bo'sh bo'lsa, yuqoridagi guruh davom etadi → limitdagi BO'LIM.
- Miqdor yoki narx bo'lmagan qatorlar (masalan «Мария хисоблаяпти», narxsiz «фанера») O'TKAZIB
  yuboriladi va ro'yxatda ko'rsatiladi.
- Natija: LimitChangeRequest (status=draft, qoralama) + qatorlar. PTO/admin limit jadvalida
  qoralamani ko'rib, tahrirlab, zanjirga yuboradi yoki admin to'g'ridan qo'llaydi.

Ishlatish:
  python manage.py limit_excel_qoralama --obyekt SHR-1 --fayl "/home/.../shirin_material.xlsx"           # ko'rish
  python manage.py limit_excel_qoralama --obyekt SHR-1 --fayl "..." --qolla                              # qoralama yaratish
  qo'shimcha: --tur material|labor|machinery|other (sukut material), --bolim "Nomi" (hamma qatorga bitta bo'lim)
"""
from decimal import Decimal, InvalidOperation

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction


def _dec(v):
    if v in (None, ""):
        return None
    try:
        return Decimal(str(v).replace(" ", "").replace("\xa0", "").replace(",", "."))
    except InvalidOperation:
        return None


class Command(BaseCommand):
    help = "Excel material ro'yxatini obyekt umumiy limitiga QORALAMA (draft) sifatida kiritadi."

    def add_arguments(self, parser):
        parser.add_argument("--obyekt", required=True, help="Obyekt kodi (masalan SHR-1)")
        parser.add_argument("--fayl", required=True, help="Excel fayl yo'li (.xlsx)")
        parser.add_argument("--varaq", default="", help="Varaq nomi (sukut — birinchi)")
        parser.add_argument("--tur", default="material", choices=["material", "labor", "machinery", "other"])
        parser.add_argument("--bolim", default="", help="Barcha qatorlarga bitta bo'lim nomi (bo'lmasa 2-ustundan)")
        parser.add_argument("--qolla", action="store_true", help="Qoralamani bazaga yozish (aks holda faqat ko'rsatadi)")
        parser.add_argument("--ustiga", action="store_true", help="Mavjud qoralama bo'lsa uni almashtirish")

    def handle(self, *args, **o):
        import openpyxl
        from projects.models import LimitChangeItem, LimitChangeRequest, Project
        from django.contrib.auth import get_user_model
        from django.conf import settings

        out = self.stdout.write
        p = Project.objects.filter(code=o["obyekt"]).first()
        if p is None:
            qid = list(Project.objects.values_list("code", "name"))
            raise CommandError(f"Obyekt topilmadi: {o['obyekt']}. Mavjud: {qid[:30]}")
        wb = openpyxl.load_workbook(o["fayl"], data_only=True)
        ws = wb[o["varaq"]] if o["varaq"] else wb.worksheets[0]

        # Sarlavha qatorini topish
        hdr, kol = None, {}
        for r in range(1, min(ws.max_row, 15) + 1):
            vals = [str(ws.cell(r, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
            if any("наимен" in v for v in vals) and any("цен" in v for v in vals):
                hdr = r
                for c, v in enumerate(vals, start=1):
                    if "наимен" in v: kol["nom"] = c
                    elif v.startswith("ед"): kol["birl"] = c
                    elif "числ" in v or "кол" in v or "миқд" in v or "miqd" in v: kol["qty"] = c
                    elif "цен" in v or "нарх" in v or "narx" in v: kol["narx"] = c
                    elif "прем" in v or "прим" in v or "изоҳ" in v or "izoh" in v: kol["izoh"] = c
                break
        if hdr is None or not all(k in kol for k in ("nom", "qty", "narx")):
            raise CommandError(f"Sarlavha topilmadi (наименование / число / цена kerak). Topilgan: {kol}")
        bolim_kol = 2 if kol["nom"] >= 3 else None
        out(f"Obyekt: {p.code} — {p.name} | varaq «{ws.title}» | sarlavha {hdr}-qator | ustunlar {kol}")

        qatorlar, otkazildi, bolim = [], [], (o["bolim"] or "").strip()
        for r in range(hdr + 1, ws.max_row + 1):
            nom = str(ws.cell(r, kol["nom"]).value or "").strip()
            if bolim_kol and not o["bolim"]:
                b = str(ws.cell(r, bolim_kol).value or "").strip()
                if b:
                    bolim = b[:200]
            if not nom:
                continue
            qty, narx = _dec(ws.cell(r, kol["qty"]).value), _dec(ws.cell(r, kol["narx"]).value)
            birl = str(ws.cell(r, kol["birl"]).value or "").strip()[:32] if "birl" in kol else ""
            izoh = str(ws.cell(r, kol["izoh"]).value or "").strip()[:500] if "izoh" in kol else ""
            if qty is None or narx is None or qty <= 0 or narx <= 0:
                otkazildi.append((r, bolim, nom, qty, narx))
                continue
            qatorlar.append({"kind": o["tur"], "name": nom[:255], "unit": birl, "quantity": qty.quantize(Decimal("0.001")),
                             "unit_price": narx.quantize(Decimal("0.01")), "note": izoh, "bolim": bolim, "masul": ""})
        jami = sum((q["quantity"] * q["unit_price"] for q in qatorlar), Decimal("0"))
        out(f"\n== Kiritiladigan qatorlar: {len(qatorlar)} ta, jami {jami:,.2f} ==".replace(",", " "))
        joriy_b = None
        for q in qatorlar:
            if q["bolim"] != joriy_b:
                joriy_b = q["bolim"]; out(f"  [{joriy_b or 'Bo‘limsiz'}]")
            out(f"    {q['name'][:40]:40} {q['quantity']} {q['unit']} × {q['unit_price']} = {q['quantity'] * q['unit_price']:,.0f}  {q['note'][:30]}".replace(",", " "))
        if otkazildi:
            out(self.style.WARNING(f"\n== O'tkazib yuborildi (miqdor/narx yo'q): {len(otkazildi)} ta =="))
            for r, b, nom, qty, narx in otkazildi:
                out(f"    {r}-qator [{b}] {nom[:40]} | miqdor={qty} narx={narx}")
        if not qolla_bool(o):
            out(self.style.WARNING("\nDRY-RUN — bazaga yozilmadi. Qoralama yaratish uchun --qolla."))
            return
        if not qatorlar:
            raise CommandError("Kiritiladigan qator yo'q.")
        lim = p.joriy_limit()
        mavjud = lim.requests.filter(status="draft")
        if mavjud.exists() and not o["ustiga"]:
            raise CommandError("Bu limitda qoralama allaqachon bor — --ustiga bilan almashtiring yoki jadvalda yakunlang.")
        adm = get_user_model().objects.filter(username=getattr(settings, "ASOSIY_ADMIN_LOGIN", "admin")).first()
        with transaction.atomic():
            mavjud.delete()
            sums = {"material": Decimal("0"), "labor": Decimal("0"), "machinery": Decimal("0"), "other": Decimal("0")}
            for q in qatorlar:
                sums[q["kind"]] += q["quantity"] * q["unit_price"]
            # qoralama = taklif: mavjud qatorlar + yangi qatorlar (mavjudlari asl bilan saqlanadi)
            req = LimitChangeRequest.objects.create(
                project=p, limit=lim,
                old_material=p.limit_material, old_labor=p.limit_labor, old_machinery=p.limit_machinery, old_other=p.limit_other,
                new_material=(p.limit_material or 0) + sums["material"], new_labor=(p.limit_labor or 0) + sums["labor"],
                new_machinery=(p.limit_machinery or 0) + sums["machinery"], new_other=(p.limit_other or 0) + sums["other"],
                reason=f"Excel'dan qoralama: {o['fayl'].split('/')[-1].split(chr(92))[-1]} ({len(qatorlar)} qator)"[:500],
                requested_by=adm, status="draft")
            items = [LimitChangeItem(request=req, kind=li.kind, name=li.name, unit=li.unit, quantity=li.quantity,
                                     unit_price=li.unit_price, note=li.note, bolim=li.bolim, masul=li.masul, asl=li,
                                     boshlangich_qty=li.boshlangich_qty, boshlangich_sum=li.boshlangich_sum)
                     for li in lim.items.all().order_by("id")]
            items += [LimitChangeItem(request=req, **q) for q in qatorlar]
            LimitChangeItem.objects.bulk_create(items)
        out(self.style.SUCCESS(f"\nQORALAMA yaratildi (#{req.pk}): {len(qatorlar)} yangi qator + {len(items) - len(qatorlar)} mavjud. "
                               f"Limit jadvali → «Қоралама» → tekshirib «Лимитни юбориш» (zanjir) yoki admin «Лимитни қўллаш»."))


def qolla_bool(o):
    return bool(o.get("qolla"))
